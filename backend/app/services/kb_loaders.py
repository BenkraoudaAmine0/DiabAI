"""Knowledge-base file loaders (LangChain BaseLoader): text, tables, images.

Every loader yields LangChain Documents of one of three content types
(metadata["content_type"]):

  text    running text of a page (PDF), chapter (EPUB) or section (Word)
  table   one table as Markdown with its caption, kept apart from the text so
          the splitter never cuts it mid-row (see kb_ingest_service.split)
  figure  what Gemini reads in an image — chart values, diagram labels,
          a photographed table — plus a short summary

Formats:
  .pdf    PyMuPDF — page text without the table areas, ruled tables
          (page.find_tables), embedded images, and scanned pages (no text
          layer) transcribed by Gemini. Page numbers, printed page labels,
          chapters from the bookmarks, title / author from the metadata.
  .docx   python-docx — paragraphs, tables and inline pictures in document
          order; each "Heading 1" starts a new section (cited as the chapter)
  .epub   ebooklib + BeautifulSoup — one Document per chapter, <table>s and
          <img>s pulled out like in PDFs
  .txt / .md   plain text (no images; Markdown tables stay as written)

Images go through `read_image(png, prompt)` (Gemini, see kb_ingest_service)
at most `max_images` times per file: anything beyond that, or any image the
model can't read, is counted in Extraction.images_skipped — never fatal.
"""

import hashlib
import io
import posixpath
import re
import warnings
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterator, Optional

from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document

# (png bytes, prompt) -> text, or "NONE" / "" for a decorative image
ReadImage = Callable[[bytes, str], str]

MIN_PAGE_TEXT = 40  # a PDF page with less text than this may be a scan
SCAN_COVERAGE = 0.4  # …if images cover at least this share of the page
MAX_IMAGE_SIDE = 1600  # px — larger renders cost tokens without helping
MAX_SCAN_SIDE = 2000
CAPTION_RE = re.compile(
    r"^\s*(table|tableau|tab\.|figure|fig\.|chart|graph|graphique|schéma|"
    r"schema|diagram|illustration|photo)\b",
    re.IGNORECASE,
)

FIGURE_PROMPT = (
    "This image comes from {where} of a medical document that is being "
    "indexed for a diabetes question-answering assistant.{caption}\n"
    "Describe it so that it can be found and cited:\n"
    "- transcribe every visible text, number, unit, axis label and legend;\n"
    "- if it is a table or a chart, give its values (tables as Markdown);\n"
    "- then sum up what it shows in 1-3 sentences.\n"
    "Write in the language of the text in the image (English if there is "
    "none). Output only the description. If the image carries no "
    "information (logo, decoration, blank area), output exactly: NONE"
)

SCAN_PROMPT = (
    "This is a scanned page of a medical document. Transcribe all of its text "
    "in reading order, exactly as written and in the same language. Render "
    "tables as Markdown tables. For a figure or photo, write one line in "
    "square brackets, e.g. [Figure: …]. Output only the page content, without "
    "commentary. If the page is blank, output exactly: NONE"
)


class IngestError(ValueError):
    """A problem with the file itself — the message is shown to the admin."""


@dataclass
class Extraction:
    tables: int = 0
    figures: int = 0  # images Gemini described
    ocr_pages: int = 0  # scanned pages Gemini transcribed
    images_skipped: int = 0  # over the per-document limit, or unreadable

    def as_dict(self) -> dict:
        return asdict(self)


@dataclass
class LoadedFile:
    docs: list[Document]
    title: Optional[str] = None  # from the file's own metadata, if any
    author: Optional[str] = None
    page_count: Optional[int] = None  # PDF only
    extraction: Extraction = field(default_factory=Extraction)


class ImageReader:
    """The per-document budget around `read_image`."""

    def __init__(
        self, read: Optional[ReadImage], limit: int, stats: Extraction
    ) -> None:
        self.read = read
        self.limit = limit if read else 0
        self.used = 0
        self.stats = stats
        self.seen: set[str] = set()  # same logo on every page → read once

    def first_time(self, data: bytes) -> bool:
        digest = hashlib.sha1(data).hexdigest()
        if digest in self.seen:
            return False
        self.seen.add(digest)
        return True

    def __call__(self, png: bytes, prompt: str) -> Optional[str]:
        if self.used >= self.limit:
            self.stats.images_skipped += 1
            return None
        self.used += 1
        try:
            text = (self.read(png, prompt) or "").strip()  # type: ignore[misc]
        except Exception as exc:  # one bad image never fails the document
            print(f"[kb] image skipped — {exc.__class__.__name__}: {exc}")
            self.stats.images_skipped += 1
            return None
        if not text or text.strip(" .`*").upper() == "NONE":
            return None  # decorative
        return text


# ── shared helpers ───────────────────────────────────────────────────────────
def clean_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def rows_to_markdown(rows: list[list[Any]]) -> Optional[str]:
    """Table rows → Markdown (first row = header). None when it isn't really
    a table: fewer than 2 non-empty rows or columns."""
    cells = [
        [" ".join(str(c or "").split()).replace("|", "\\|") for c in row]
        for row in rows
    ]
    cells = [r for r in cells if any(r)]
    if len(cells) < 2:
        return None
    width = max(len(r) for r in cells)
    cells = [r + [""] * (width - len(r)) for r in cells]
    keep = [i for i in range(width) if any(r[i] for r in cells)]
    if len(keep) < 2:
        return None
    cells = [[r[i] for i in keep] for r in cells]
    header, *body = cells
    lines = ["| " + " | ".join(header) + " |", "|" + " --- |" * len(header)]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _labelled(kind: str, caption: Optional[str], body: str) -> str:
    return f"{kind}: {caption}\n{body}" if caption else f"{kind}\n{body}"


def _figure_prompt(where: str, caption: Optional[str]) -> str:
    return FIGURE_PROMPT.format(
        where=where, caption=f' Its caption reads: "{caption}".' if caption else ""
    )


def _to_png(data: bytes, min_px: int) -> Optional[bytes]:
    """Any Pillow-readable image → PNG (Gemini doesn't take GIF/BMP/TIFF/
    EMF…). None when unreadable or smaller than min_px on a side."""
    from PIL import Image

    try:
        with Image.open(io.BytesIO(data)) as img:
            if min(img.size) < min_px:
                return None
            img = img.convert("RGBA" if "A" in img.getbands() else "RGB")
            img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
            out = io.BytesIO()
            img.save(out, format="PNG")
            return out.getvalue()
    except Exception:
        return None


# ── PDF ──────────────────────────────────────────────────────────────────────
def _chapter_at(chapters: list[tuple[int, str]], page_index: int) -> Optional[str]:
    current = None
    for start, title in chapters:
        if start > page_index:
            break
        current = title
    return current


def _overlap(a: Any, b: Any) -> float:
    """Share of rect a's area that lies inside rect b."""
    inter = a & b
    if inter.is_empty or a.is_empty:
        return 0.0
    return (inter.width * inter.height) / (a.width * a.height)


def _pdf_caption(blocks: list, rect: Any, pymupdf: Any) -> Optional[str]:
    """The text block just above or below a table / image that looks like
    its caption ("Table 2. …", "Figure 3 – …")."""
    best: Optional[tuple[float, str]] = None
    for b in blocks:
        text = " ".join(b[4].split())
        if not text or not CAPTION_RE.match(text):
            continue
        br = pymupdf.Rect(b[:4])
        gap = min(abs(rect.y0 - br.y1), abs(br.y0 - rect.y1))
        if gap <= 60 and (best is None or gap < best[0]):
            best = (gap, text[:200])
    return best[1] if best else None


def _render(page: Any, rect: Any, max_side: int) -> bytes:
    zoom = min(200 / 72, max_side / max(rect.width, rect.height, 1))
    import pymupdf

    return page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=rect).tobytes("png")


class PdfLoader(BaseLoader):
    """One `text` Document per page plus `table` / `figure` Documents.
    After load(): title, author, page_count."""

    def __init__(self, file_path: str, reader: ImageReader, min_image_px: int):
        self.file_path = file_path
        self.reader = reader
        self.min_image_px = min_image_px
        self.title: Optional[str] = None
        self.author: Optional[str] = None
        self.page_count: Optional[int] = None

    def lazy_load(self) -> Iterator[Document]:
        import pymupdf

        getattr(pymupdf, "no_recommend_layout", lambda: None)()  # log noise
        try:
            pdf = pymupdf.open(self.file_path)
        except Exception as exc:
            raise IngestError(f"Could not read this PDF: {exc}") from exc
        try:
            if pdf.needs_pass:
                raise IngestError("This PDF is password-protected.")
            meta = pdf.metadata or {}
            self.title = meta.get("title") or None
            self.author = meta.get("author") or None
            self.page_count = pdf.page_count
            try:
                chapters = sorted(
                    (entry[2] - 1, " ".join(str(entry[1]).split()))
                    for entry in pdf.get_toc(simple=True)
                    if entry[0] == 1 and entry[2] >= 1 and str(entry[1]).strip()
                )
            except Exception:
                chapters = []  # broken outline — citations just lack chapters
            for page in pdf:
                label = (page.get_label() or "").strip()
                base = {
                    "page": page.number + 1,
                    # printed page number, when it differs from the PDF page
                    "page_label": label if label and label != str(page.number + 1) else None,
                    "chapter": _chapter_at(chapters, page.number),
                }
                yield from self._page(page, base, pymupdf)
        finally:
            pdf.close()

    def _page(self, page: Any, base: dict, pymupdf: Any) -> Iterator[Document]:
        stats = self.reader.stats
        blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
        out: list[Document] = []

        table_rects = []
        try:
            tables = page.find_tables().tables
        except Exception:
            tables = []
        for table in tables:
            md = rows_to_markdown(table.extract())
            if not md:
                continue
            rect = pymupdf.Rect(table.bbox)
            table_rects.append(rect)
            stats.tables += 1
            out.append(
                Document(
                    page_content=_labelled("Table", _pdf_caption(blocks, rect, pymupdf), md),
                    metadata={**base, "content_type": "table"},
                )
            )

        text = clean_text(
            "\n".join(
                b[4]
                for b in blocks
                if not any(_overlap(pymupdf.Rect(b[:4]), r) > 0.5 for r in table_rects)
            )
        )

        images = []
        for info in page.get_image_info(hashes=True):
            rect = pymupdf.Rect(info["bbox"]) & page.rect
            if rect.is_empty or min(info["width"], info["height"]) < self.min_image_px:
                continue
            if rect.width < 40 or rect.height < 40:  # points: thumbnails, bullets
                continue
            images.append((rect, info.get("digest") or b""))

        page_area = page.rect.width * page.rect.height
        coverage = sum(r.width * r.height for r, _ in images) / page_area
        if len(text) < MIN_PAGE_TEXT and coverage >= SCAN_COVERAGE:
            # a scanned page: the whole page is the image — transcribe it
            ocr = self.reader(_render(page, page.rect, MAX_SCAN_SIDE), SCAN_PROMPT)
            if ocr:
                stats.ocr_pages += 1
                text = clean_text(f"{text}\n{ocr}")
                base = {**base, "ocr": True}
        else:
            where = f"page {base['page']}"
            for rect, digest in images:
                if digest and not self.reader.first_time(digest):
                    continue
                caption = _pdf_caption(blocks, rect, pymupdf)
                description = self.reader(
                    _render(page, rect, MAX_IMAGE_SIDE), _figure_prompt(where, caption)
                )
                if description:
                    stats.figures += 1
                    out.append(
                        Document(
                            page_content=_labelled("Figure", caption, description),
                            metadata={**base, "content_type": "figure"},
                        )
                    )

        if text:
            yield Document(page_content=text, metadata={**base, "content_type": "text"})
        yield from out


# ── Word (.docx) ─────────────────────────────────────────────────────────────
def _is_section_heading(paragraph: Any) -> bool:
    try:
        name = (paragraph.style.name or "").lower()
    except Exception:
        return False
    return name in ("heading 1", "title", "titre 1", "titre")


class DocxLoader(BaseLoader):
    """Word documents in reading order: one `text` Document per "Heading 1"
    section (the heading becomes the cited chapter), `table` Documents for
    tables and `figure` Documents for pictures."""

    def __init__(self, file_path: str, reader: ImageReader, min_image_px: int):
        self.file_path = file_path
        self.reader = reader
        self.min_image_px = min_image_px
        self.title: Optional[str] = None
        self.author: Optional[str] = None

    def lazy_load(self) -> Iterator[Document]:
        import docx
        from docx.table import Table

        try:
            word = docx.Document(self.file_path)
        except Exception as exc:
            raise IngestError(f"Could not read this Word document: {exc}") from exc
        props = word.core_properties
        self.title = props.title or None
        self.author = props.author or None

        blocks = list(word.iter_inner_content())
        texts = [getattr(b, "text", "").strip() for b in blocks]

        def caption_near(i: int) -> Optional[str]:
            for j in (i - 1, i + 1):
                if 0 <= j < len(texts) and CAPTION_RE.match(texts[j]):
                    return texts[j][:200]
            return None

        stats = self.reader.stats
        section: Optional[str] = None
        lines: list[str] = []
        side: list[Document] = []  # tables / figures of the current section

        def flush() -> Iterator[Document]:
            text = clean_text("\n".join(lines))
            if text:
                yield Document(
                    page_content=text,
                    metadata={"chapter": section, "content_type": "text"},
                )
            yield from side
            lines.clear()
            side.clear()

        for i, block in enumerate(blocks):
            if isinstance(block, Table):
                try:
                    rows = [[cell.text for cell in row.cells] for row in block.rows]
                except Exception:
                    rows = []
                md = rows_to_markdown(rows)
                if md:
                    stats.tables += 1
                    side.append(
                        Document(
                            page_content=_labelled("Table", caption_near(i), md),
                            metadata={"chapter": section, "content_type": "table"},
                        )
                    )
                continue

            if _is_section_heading(block) and texts[i]:
                yield from flush()
                section = texts[i][:200]
            if texts[i]:
                lines.append(texts[i])

            for rid in block._element.xpath(".//a:blip/@r:embed"):
                part = word.part.related_parts.get(rid)
                blob = getattr(part, "blob", None)
                if not blob or not self.reader.first_time(blob):
                    continue
                png = _to_png(blob, self.min_image_px)
                if png is None:
                    continue
                caption = caption_near(i)
                where = f"the section “{section}”" if section else "a page"
                description = self.reader(png, _figure_prompt(where, caption))
                if description:
                    stats.figures += 1
                    side.append(
                        Document(
                            page_content=_labelled("Figure", caption, description),
                            metadata={"chapter": section, "content_type": "figure"},
                        )
                    )
        yield from flush()


# ── EPUB ─────────────────────────────────────────────────────────────────────
def _epub_meta(book: Any, name: str) -> Optional[str]:
    try:
        values = book.get_metadata("DC", name)
        return (str(values[0][0]).strip() or None) if values else None
    except Exception:
        return None


def _epub_toc_titles(toc: Any) -> dict[str, str]:
    """Chapter file → title, from the EPUB table of contents.

    Several entries can point into one file, so they are ranked: an entry
    pointing at the file itself beats one pointing at an anchor inside it
    (the chapter, not its sub-sections), then a plain link beats a section
    with children (the chapter, not "Part II" that starts at it)."""
    ranked: dict[str, tuple[int, str]] = {}

    def add(entry: Any, is_section: bool) -> None:
        href = getattr(entry, "href", None)
        title = getattr(entry, "title", None)
        if not href or not title:
            return
        file = href.split("#")[0]
        rank = (2 if "#" in href else 0) + (1 if is_section else 0)
        if file not in ranked or rank < ranked[file][0]:
            ranked[file] = (rank, str(title).strip())

    def walk(entries: Any) -> None:
        for entry in entries:
            if isinstance(entry, tuple):  # (Section, [children])
                add(entry[0], is_section=True)
                walk(entry[1])
            else:
                add(entry, is_section=False)

    walk(toc or [])
    return {file: title for file, (_rank, title) in ranked.items()}


def _html_caption(tag: Any) -> Optional[str]:
    cap = tag.find("caption") or (
        tag.find_parent("figure").find("figcaption") if tag.find_parent("figure") else None
    )
    if cap is not None:
        text = cap.get_text(" ", strip=True)
        if text:
            return text[:200]
    prev = tag.find_previous(["p", "h1", "h2", "h3", "h4", "div"])
    if prev is not None:
        text = prev.get_text(" ", strip=True)
        if CAPTION_RE.match(text):
            return text[:200]
    return None


class EpubLoader(BaseLoader):
    """LangChain loader for EPUB books: one `text` Document per chapter
    (spine item) in reading order, with metadata {"chapter": <title from the
    TOC>}, plus `table` / `figure` Documents.
    After load(), `title` and `author` hold the book's own metadata."""

    def __init__(
        self, file_path: str, reader: Optional[ImageReader] = None, min_image_px: int = 120
    ):
        self.file_path = file_path
        self.reader = reader or ImageReader(None, 0, Extraction())
        self.min_image_px = min_image_px
        self.title: Optional[str] = None
        self.author: Optional[str] = None

    def lazy_load(self) -> Iterator[Document]:
        from bs4 import BeautifulSoup
        from ebooklib import ITEM_DOCUMENT, epub

        with warnings.catch_warnings():
            # ebooklib warns about a future default; we still want the NCX
            # table of contents (EPUB 2 books only have that one)
            warnings.simplefilter("ignore")
            book = epub.read_epub(self.file_path, options={"ignore_ncx": False})

        self.title = _epub_meta(book, "title")
        self.author = _epub_meta(book, "creator")
        toc = _epub_toc_titles(book.toc)
        stats = self.reader.stats

        chapter: Optional[str] = None
        for item_id, _linear in book.spine:
            item = book.get_item_with_id(item_id)
            if (
                item is None
                or item.get_type() != ITEM_DOCUMENT
                or isinstance(item, epub.EpubNav)  # EPUB 3 clickable TOC
                or "nav" in (getattr(item, "properties", None) or [])
            ):
                continue
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")  # XHTML parsed as HTML
                soup = BeautifulSoup(item.get_content(), "html.parser")
            for tag in soup(["script", "style"]):
                tag.decompose()
            if not clean_text(soup.get_text("\n")):
                continue  # cover image, blank page…
            heading = soup.find(["h1", "h2", "h3"])
            chapter = (
                toc.get(item.get_name())
                or (heading.get_text(" ", strip=True) if heading else None)
                or chapter  # continuation file of the previous chapter
            )
            meta = {"source": self.file_path, "chapter": chapter}

            side: list[Document] = []
            for table in soup.find_all("table"):
                if table.find_parent("table") is not None:
                    continue  # nested: read with its parent
                rows = [
                    [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
                    for tr in table.find_all("tr")
                ]
                md = rows_to_markdown(rows)
                if md:
                    stats.tables += 1
                    side.append(
                        Document(
                            page_content=_labelled("Table", _html_caption(table), md),
                            metadata={**meta, "content_type": "table"},
                        )
                    )
                table.decompose()

            folder = posixpath.dirname(item.get_name())
            for img in soup.find_all("img"):
                src = (img.get("src") or "").split("#")[0]
                image_item = (
                    book.get_item_with_href(posixpath.normpath(posixpath.join(folder, src)))
                    if src and not src.startswith(("http:", "https:", "data:"))
                    else None
                )
                blob = image_item.get_content() if image_item is not None else None
                if blob and self.reader.first_time(blob):
                    png = _to_png(blob, self.min_image_px)
                    if png is not None:
                        caption = _html_caption(img) or (img.get("alt") or "").strip()[:200] or None
                        where = f"the chapter “{chapter}”" if chapter else "a book"
                        description = self.reader(png, _figure_prompt(where, caption))
                        if description:
                            stats.figures += 1
                            side.append(
                                Document(
                                    page_content=_labelled("Figure", caption, description),
                                    metadata={**meta, "content_type": "figure"},
                                )
                            )
                img.decompose()

            text = clean_text(soup.get_text("\n"))
            if text:
                yield Document(page_content=text, metadata={**meta, "content_type": "text"})
            yield from side


# ── plain text ───────────────────────────────────────────────────────────────
def read_text(path: Path) -> str:
    """UTF-8 (with or without BOM), UTF-16 (BOM), else Windows-1252 — the
    legacy encoding of French text files. Not LangChain's TextLoader: its
    encoding=None means cp1252 on Windows (garbled UTF-8 accents), and its
    autodetect path crashes with chardet >= 6 (langchain-community 0.4)."""
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def load(
    path: Path,
    read_image: Optional[ReadImage] = None,
    max_images: int = 0,
    min_image_px: int = 120,
) -> LoadedFile:
    """Any supported file → LoadedFile. Raises IngestError for unreadable
    files (the message is shown to the admin)."""
    ext = path.suffix.lower()
    stats = Extraction()
    reader = ImageReader(read_image, max_images, stats)

    if ext == ".pdf":
        loader = PdfLoader(str(path), reader, min_image_px)
        docs = loader.load()
        loaded = LoadedFile(
            docs=docs, title=loader.title, author=loader.author, page_count=loader.page_count
        )
    elif ext == ".docx":
        dloader = DocxLoader(str(path), reader, min_image_px)
        loaded = LoadedFile(docs=dloader.load(), title=dloader.title, author=dloader.author)
    elif ext == ".epub":
        eloader = EpubLoader(str(path), reader, min_image_px)
        try:
            docs = eloader.load()
        except IngestError:
            raise
        except Exception as exc:
            raise IngestError(f"Could not read this EPUB: {exc}") from exc
        loaded = LoadedFile(docs=docs, title=eloader.title, author=eloader.author)
    elif ext in (".txt", ".md"):
        try:
            text = read_text(path)
        except OSError as exc:
            raise IngestError(f"Could not read this text file: {exc}") from exc
        loaded = LoadedFile(
            docs=[
                Document(
                    page_content=text,
                    metadata={"source": str(path), "content_type": "text"},
                )
            ]
        )
    else:
        raise IngestError(f"Unsupported file type '{ext}'.")

    loaded.extraction = stats
    return loaded
